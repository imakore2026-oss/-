"""
記入済みPDF フォルダを監視し、PDFフォームの入力内容を自動で
Excel（受注・作業リスト.xlsx）に一覧化するツール。

使い方は README.md を参照してください。
Ctrl+C で終了します。
"""
import json
import sys
import time
from datetime import datetime
from pathlib import Path

try:
    from pypdf import PdfReader
except ImportError:
    print("エラー: pypdf がインストールされていません。作業開始.bat から実行してください。")
    sys.exit(1)

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    print("エラー: openpyxl がインストールされていません。作業開始.bat から実行してください。")
    sys.exit(1)

TOOL_DIR = Path(__file__).resolve().parent
BASE_DIR = TOOL_DIR.parent
INBOX_DIR = BASE_DIR / "記入済みPDF"
OUTPUT_XLSX = BASE_DIR / "受注・作業リスト.xlsx"
STATE_FILE = TOOL_DIR / ".state.json"
POLL_INTERVAL_SEC = 3

EXTRA_COLUMNS = ["元ファイル名", "取込日時"]


def load_field_map(name):
    with open(TOOL_DIR / name, encoding="utf-8") as f:
        return json.load(f)


FIELD_MAPS = [
    load_field_map("field_map_order.json"),
    load_field_map("field_map_work.json"),
]


def load_state():
    if STATE_FILE.exists():
        try:
            return json.loads(STATE_FILE.read_text(encoding="utf-8"))
        except Exception:
            return {}
    return {}


def save_state(state):
    STATE_FILE.write_text(
        json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def ensure_workbook():
    if OUTPUT_XLSX.exists():
        wb = load_workbook(OUTPUT_XLSX)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    for fmap in FIELD_MAPS:
        sheet_name = fmap["_form_name"]
        headers = [f["label"] for f in fmap["_fields"]] + EXTRA_COLUMNS
        if sheet_name not in wb.sheetnames:
            ws = wb.create_sheet(sheet_name)
            ws.append(headers)
        else:
            ws = wb[sheet_name]
            if ws.max_row == 0 or [c.value for c in ws[1]] != headers:
                if ws.max_row == 0:
                    ws.append(headers)
    return wb


def checkbox_checked(value):
    return str(value) in ("/Yes", "Yes", "/On", "On")


def match_field_map(field_names):
    for fmap in FIELD_MAPS:
        if fmap["_signature_field"] in field_names:
            return fmap
    return None


def extract_row(pdf_path, fmap):
    reader = PdfReader(str(pdf_path))
    fields = reader.get_fields() or {}
    row = []
    for spec in fmap["_fields"]:
        f = fields.get(spec["name"])
        value = f.get("/V") if f else None
        if spec.get("type") == "checkbox":
            row.append("○" if checkbox_checked(value) else "")
        else:
            row.append("" if value is None else str(value))
    return row


def upsert_row(ws, filename, row_values):
    header = [c.value for c in ws[1]]
    filename_col = header.index("元ファイル名") + 1
    timestamp_col = header.index("取込日時") + 1

    target_row = None
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=filename_col).value == filename:
            target_row = r
            break
    if target_row is None:
        target_row = ws.max_row + 1

    for col, value in enumerate(row_values, start=1):
        ws.cell(row=target_row, column=col, value=value)
    ws.cell(row=target_row, column=filename_col, value=filename)
    ws.cell(
        row=target_row,
        column=timestamp_col,
        value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    )


def process_file(pdf_path, wb):
    reader = PdfReader(str(pdf_path))
    fields = reader.get_fields()
    if not fields:
        print(f"  スキップ（入力欄のないPDFです）: {pdf_path.name}")
        return True

    fmap = match_field_map(fields.keys())
    if fmap is None:
        print(f"  スキップ（対応していない様式です）: {pdf_path.name}")
        return True

    row_values = extract_row(pdf_path, fmap)
    ws = wb[fmap["_form_name"]]
    upsert_row(ws, pdf_path.name, row_values)
    print(f"  取込みました: {pdf_path.name} → シート「{fmap['_form_name']}」")
    return True


def main():
    INBOX_DIR.mkdir(exist_ok=True)
    print("=" * 60)
    print("PDF自動リスト化ツール")
    print(f"監視フォルダ : {INBOX_DIR}")
    print(f"出力Excel   : {OUTPUT_XLSX}")
    print("このウィンドウを開いたままにしておくと、")
    print("上記フォルダにPDFを保存するたびに自動でExcelへ反映されます。")
    print("終了するには Ctrl+C を押してください。")
    print("=" * 60)

    state = load_state()

    while True:
        try:
            pdf_files = sorted(INBOX_DIR.glob("*.pdf"))
            changed = []
            for pdf_path in pdf_files:
                mtime = pdf_path.stat().st_mtime
                key = pdf_path.name
                if state.get(key) != mtime:
                    changed.append((pdf_path, mtime))

            if changed:
                wb = ensure_workbook()
                for pdf_path, mtime in changed:
                    try:
                        process_file(pdf_path, wb)
                    except Exception as e:
                        print(f"  読み込み失敗（保存中の可能性）: {pdf_path.name} ({e})")
                        continue
                    try:
                        wb.save(OUTPUT_XLSX)
                    except PermissionError:
                        print(
                            "  Excelファイルが開かれているため保存できません。"
                            "Excelを閉じると自動的に再試行します。"
                        )
                        break
                    else:
                        state[pdf_path.name] = mtime
                        save_state(state)

        except Exception as e:
            print(f"エラーが発生しました（続行します）: {e}")

        time.sleep(POLL_INTERVAL_SEC)


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n終了しました。")
